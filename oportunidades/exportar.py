"""Exporta todo a data/oportunidades.xlsx: Resumen, Empresas, Vacantes, Seguimiento."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config
from .contactos import contactos, url_linkedin_empresa
from .db import ESTADOS, Store

ENCABEZADO = "B9FF66"
COLORES_ESTADO = {
    "Descubierta": "F2F2F5", "Me interesa": "FEF3C7", "Investigando": "E0F2FE", "Contacté": "EDE9FE",
    "Aplicado": "DBEAFE", "Entrevista": "D1FAE5", "Oferta": "A7F3D0", "Descartada": "E5E7EB",
}
_ENCABEZADO = PatternFill("solid", fgColor=ENCABEZADO)


def _hoja(wb: Workbook, titulo: str, columnas: list[tuple[str, int]], filas: list[list],
          col_estado: int | None = None, primera: bool = False):
    ws = wb.active if primera else wb.create_sheet()
    ws.title = titulo
    ws.append([c for c, _ in columnas])
    for i, (_, ancho) in enumerate(columnas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = Font(bold=True, color="1A1A1E"), _ENCABEZADO
    for fila in filas:
        ws.append(fila)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str) and len(cell.value) > 40)
            if isinstance(cell.value, str) and cell.value.startswith("http") and "\n" not in cell.value:
                cell.hyperlink, cell.style = cell.value, "Hyperlink"
        if col_estado is not None:
            c = row[col_estado]
            if c.value in COLORES_ESTADO:
                c.fill = PatternFill("solid", fgColor=COLORES_ESTADO[c.value])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def _enlaces(items: list[dict], clave: str) -> str:
    return "\n".join(" · ".join(x for x in (i.get(clave, ""), i.get("url", "")) if x) for i in items)


def exportar(store: Store, destino: Path | None = None) -> Path:
    destino = destino or config.xlsx_path()
    r = store.resumen()
    wb = Workbook()

    filas = [[e, r["por_estado"][e]] for e in ESTADOS]
    filas += [[], ["Empresas", r["total_empresas"]], ["Contactadas o aplicadas esta semana", r["contactadas_semana"]],
              ["Empresas con vacantes abiertas", r["empresas_con_vacantes"]],
              ["Vacantes abiertas", r["vacantes_activas"]],
              ["Último rastreo de vacantes", (r["ultima_corrida"] or {}).get("fin", "")]]
    if r["seguimientos"]:
        filas += [[], ["Seguimientos próximos", ""]]
        filas += [[f"{s['seguimiento']} · {s['estado']}", s["nombre"]] for s in r["seguimientos"]]
    _hoja(wb, "Resumen", [("Estado", 36), ("Cantidad", 40)], filas, col_estado=0, primera=True)

    filas = []
    for e in store.empresas():
        cts = contactos(e)
        filas.append([
            e["nombre"], e["tipo"], e["industria"], e["estado"], e["presencia_rd"],
            _enlaces(e["fuentes_presencia"], "titulo"), "\n".join(e["equipos"]),
            "\n".join(f"{x['rol']}: {x['razon']}".rstrip(": ") for x in e["roles"]),
            _enlaces(e["evidencia"], "texto"), "\n".join(e["razones"]), "\n".join(e["alertas"]),
            e["sitio_web"], e["url_empleos"], e["portal_rd"], e["talento"], url_linkedin_empresa(e),
            "; ".join(c["rol"] for c in cts), "\n".join(c["linkedin"] for c in cts),
            e["vacantes"], e["vacantes_activas"], e["nota_rastreo"], e["nota"], e["seguimiento"],
            e["mensaje_path"], e["carta_path"], e["creado"][:10]])
    _hoja(wb, "Empresas", [
        ("Empresa", 26), ("Tipo", 18), ("Industria", 18), ("Estado", 14), ("Presencia en RD", 40),
        ("Fuentes de presencia", 50), ("Equipos que encajan", 34), ("Roles probables", 50), ("Evidencia", 50),
        ("Por qué encaja", 50), ("Señales de alerta", 44), ("Sitio web", 28), ("Página de empleos", 32),
        ("Portal filtrado a RD", 32), ("Talento / postulación general", 32), ("LinkedIn de la empresa", 32),
        ("Roles a contactar", 36), ("Búsqueda LinkedIn", 56), ("Vacantes", 16), ("Vacantes abiertas", 10),
        ("Nota de rastreo", 34), ("Notas", 30), ("Seguimiento", 12), ("Mensaje", 30), ("Carta", 30),
        ("Descubierta", 12)], filas, col_estado=3)

    filas = [[v["id"], v["empresa_nombre"], v["titulo"], v["ubicacion"], v["url"], v["fecha_publicacion"],
              v["fuente"], "sí" if v["activa"] else "ya no aparece", v["primera_vez"][:10]]
             for v in store.vacantes()]
    _hoja(wb, "Vacantes", [("ID", 6), ("Empresa", 26), ("Puesto", 40), ("Ubicación", 22), ("Enlace", 44),
                           ("Publicada", 12), ("Fuente", 16), ("Activa", 14), ("Encontrada", 12)], filas)

    filas = [[h["fecha"].replace("T", " "), h["empresa_nombre"], h["de_estado"], h["a_estado"], h["nota"]]
             for h in store.historial()]
    _hoja(wb, "Seguimiento", [("Fecha", 20), ("Empresa", 28), ("De", 14), ("A", 14), ("Nota", 50)], filas,
          col_estado=3)

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


def exportar_seguro(store: Store) -> str:
    """Para usar después de cada cambio: si el Excel está abierto (Windows lo bloquea), avisa y sigue."""
    try:
        return str(exportar(store))
    except PermissionError:
        return "No se pudo actualizar oportunidades.xlsx (¿está abierto en Excel?). Ciérralo y corre scripts/exportar.py."
