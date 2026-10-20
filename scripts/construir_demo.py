#!/usr/bin/env python3
"""Construye la demo pública (archivos estáticos, sin Python ni servidor) para GitHub Pages.

Uso:
    python scripts/construir_demo.py               # escribe en docs/
    python scripts/construir_demo.py --destino X   # en otra carpeta

Solo usa datos inventados: las empresas de scripts/demo.py y el perfil de ejemplo (profile.example.yaml),
creados en una carpeta temporal. Nunca lee tu perfil ni tu base de datos. La página es la misma del
tablero local (templates/index.html con publica=True); app.js lee las respuestas guardadas en
demo-datos.js en lugar de llamar a la API, y no guarda nada.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import shutil
import tempfile
import zipfile
from pathlib import Path

from _comun import ROOT

DESTINO = ROOT / "docs"
# Lo que la página pide al cargar (más el detalle de cada empresa, que se agrega abajo).
RUTAS = ["/api/resumen", "/api/empresas", "/api/vacantes", "/api/perfil", "/api/buscar"]


def _sin_rutas(obj):
    """Quita las rutas de archivos locales (mensaje_path, carta_path…): no sirven en la web y revelan carpetas."""
    if isinstance(obj, dict):
        return {k: _sin_rutas(v) for k, v in obj.items() if not k.endswith("_path")}
    if isinstance(obj, list):
        return [_sin_rutas(v) for v in obj]
    return obj


def _excel_sin_rutas(origen: Path, destino: Path, carpeta_datos: Path) -> str:
    """Copia el Excel cambiando las rutas absolutas de los documentos por data/documentos/…; devuelve su texto."""
    from openpyxl import load_workbook

    libro = load_workbook(origen)
    prefijos = {str(carpeta_datos.resolve()), str(carpeta_datos)}  # en Windows, la ruta corta (ALEJAN~1) y la larga
    for hoja in libro.worksheets:
        for fila in hoja.iter_rows():
            for celda in fila:
                for prefijo in prefijos:
                    if isinstance(celda.value, str) and prefijo in celda.value:
                        celda.value = celda.value.replace(prefijo, "data").replace("\\", "/")
    libro.save(destino)
    with zipfile.ZipFile(destino) as z:
        return "".join(z.read(n).decode("utf-8", "replace") for n in z.namelist() if n.endswith(".xml"))


def construir(destino: Path = DESTINO) -> list[Path]:
    from oportunidades import config

    tmp = Path(tempfile.mkdtemp(prefix="oportunidades-demo-"))
    os.environ[config.DATA_ENV] = str(tmp / "data")
    os.environ[config.PROFILE_ENV] = str(tmp / "profile.yaml")
    try:
        import demo
        from flask import render_template

        import app as tablero
        from oportunidades import perfil as pf

        xlsx = demo.crear_datos(tmp / "data")
        pf.guardar(pf.importar(config.ROOT / "profile.example.yaml"))
        web = tablero.crear_app(demo=True)
        cliente = web.test_client()
        datos = {r: cliente.get(r).get_json() for r in RUTAS}
        for e in datos["/api/empresas"]:
            datos[f"/api/empresas/{e['id']}"] = cliente.get(f"/api/empresas/{e['id']}").get_json()
        datos = _sin_rutas(datos)
        with web.test_request_context():
            html = render_template("index.html", **tablero.datos_pagina(publica=True))

        # DEMO_FECHA: el día de la construcción. app.js mueve todas las fechas a partir de hoy al cargar la página.
        js = ("// Datos INVENTADOS para la demo pública. Generado por scripts/construir_demo.py; no lo edites a mano.\n"
              f'window.DEMO_FECHA = "{dt.date.today().isoformat()}";\n'
              f"window.DEMO = {json.dumps(datos, ensure_ascii=False, indent=1)};\n")
        excel = tmp / "oportunidades-demo.xlsx"
        for texto in (js, html, _excel_sin_rutas(xlsx, excel, tmp / "data")):
            for prohibido in {str(tmp), str(tmp.resolve()), str(Path.home()), str(ROOT)}:
                if prohibido in texto or prohibido.replace("\\", "\\\\") in texto:
                    raise RuntimeError(f"La demo incluiría una ruta local ({prohibido}); no se escribió nada.")

        destino.mkdir(parents=True, exist_ok=True)
        salida = {"index.html": html, "demo-datos.js": js, ".nojekyll": ""}
        for nombre, texto in salida.items():
            (destino / nombre).write_text(texto, encoding="utf-8", newline="\n")
        for nombre in ("style.css", "app.js", "sugerencias.js", "esco.json"):
            shutil.copyfile(ROOT / "static" / nombre, destino / nombre)  # esco.json: python scripts/construir_esco.py
        shutil.copyfile(excel, destino / "oportunidades-demo.xlsx")
        return [destino / n for n in (*salida, "style.css", "app.js", "sugerencias.js", "esco.json", "oportunidades-demo.xlsx")]
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Construye la demo pública estática (GitHub Pages)")
    ap.add_argument("--destino", type=Path, default=DESTINO)
    args = ap.parse_args(argv)
    for p in construir(args.destino):
        print(p)
    print("Abre index.html en esa carpeta para verla. En GitHub: Settings → Pages → rama main, carpeta /docs.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
