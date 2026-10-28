from openpyxl import load_workbook

from conftest import empresa
from oportunidades import config
from oportunidades.db import Store
from oportunidades.exportar import COLORES_ESTADO, exportar, exportar_seguro


def _datos(store):
    e = empresa(store, "Banco Ejemplo", tipo="grupo local", presencia_rd="Oficinas en Santo Domingo",
                equipos=["Datos y analítica"], roles=[{"rol": "BI Analyst", "razon": "Tableau"}],
                razones=["a", "b", "c"], alertas=["Solo presencial"], roles_contacto=["Head of Data"],
                url_empleos="https://b.do/empleos")
    store.upsert_vacante("Analista de Datos", e["id"], url="https://b.do/empleos/1")
    store.cambiar_estado(e["id"], "Me interesa")
    store.cambiar_estado(e["id"], "Contacté")
    empresa(store, "Telco Ejemplo")


def test_xlsx_con_hojas_filtros_encabezados_fijos_y_colores(store):
    _datos(store)
    path = exportar(store)
    assert path == config.xlsx_path() and path.exists()
    wb = load_workbook(path)
    assert wb.sheetnames == ["Resumen", "Empresas", "Vacantes", "Seguimiento"]
    for ws in wb.worksheets:
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref
        assert ws["A1"].font.bold
    encabezados = [c.value for c in wb["Empresas"][1]]
    assert "Ajuste" not in encabezados and "Interés" not in encabezados
    emp = {r[0].value: r for r in wb["Empresas"].iter_rows(min_row=2)}
    banco = emp["Banco Ejemplo"]
    assert banco[3].value == "Contacté" and banco[3].fill.fgColor.rgb.endswith(COLORES_ESTADO["Contacté"])
    fila = dict(zip(encabezados, [c.value for c in banco]))
    assert "ejemplo.do/contacto" in fila["Fuentes de presencia"] and fila["Equipos que encajan"] == "Datos y analítica"
    assert fila["Roles probables"] == "BI Analyst: Tableau" and fila["Vacantes"] == "hay vacantes"
    assert "Head of Data" in fila["Cargos a contactar"] and "linkedin.com/search" in fila["LinkedIn"]
    vac = list(wb["Vacantes"].iter_rows(min_row=2, values_only=True))
    assert vac[0][2] == "Analista de Datos" and len(vac[0]) == 9
    seg = [(r[2].value, r[3].value) for r in wb["Seguimiento"].iter_rows(min_row=2)]
    assert ("Me interesa", "Contacté") in seg and ("Descubierta", "Me interesa") in seg
    resumen = {r[0]: r[1] for r in wb["Resumen"].iter_rows(min_row=2, values_only=True) if r[0]}
    assert resumen["Contacté"] == 1 and resumen["Descubierta"] == 1 and resumen["Contactadas o aplicadas esta semana"] == 1


def test_se_regenera_despues_de_cada_cambio():
    with Store(after_change=lambda s: exportar_seguro(s)) as s:
        e = empresa(s, "Telco Ejemplo")
        assert config.xlsx_path().exists()
        s.cambiar_estado(e["id"], "Me interesa")
        estados = [r[3] for r in load_workbook(config.xlsx_path())["Empresas"].iter_rows(min_row=2, values_only=True)]
        assert estados == ["Me interesa"]


def test_excel_abierto_no_rompe(store, monkeypatch):
    import oportunidades.exportar as ex

    def bloqueado(*a, **k):
        raise PermissionError("abierto")
    monkeypatch.setattr(ex, "exportar", bloqueado)
    assert "abierto en Excel" in ex.exportar_seguro(store)
